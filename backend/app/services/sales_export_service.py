from datetime import datetime
from io import BytesIO
from typing import Any, Dict, List
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.config import settings

MONTH_NAMES_ES = (
    "",
    "Enero",
    "Febrero",
    "Marzo",
    "Abril",
    "Mayo",
    "Junio",
    "Julio",
    "Agosto",
    "Septiembre",
    "Octubre",
    "Noviembre",
    "Diciembre",
)

CURRENCY_FORMAT = '"$"#,##0.00'
DATE_FORMAT = "dd/mm/yyyy hh:mm"
HEADER_FILL = PatternFill("solid", fgColor="312E81")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _write_text(cell: Any, value: Any) -> None:
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"


def _to_excel_datetime(value: Any) -> datetime | None:
    if not value:
        return None
    parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    return parsed.replace(tzinfo=None)


def build_sales_workbook(
    rows: List[Dict[str, Any]],
    year: int,
    month: int,
    generated_at: datetime | None = None,
) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    details = workbook.create_sheet("Detalle de ventas")
    period_label = f"{MONTH_NAMES_ES[month]} {year}"

    completed_sales: Dict[str, Dict[str, Any]] = {}
    cancelled_sales: set[str] = set()
    completed_units = 0.0

    for row in rows:
        sale_id = str(row["sale_id"])
        if row["status"] == "COMPLETED":
            completed_sales[sale_id] = row
            if row["product_id"] is not None:
                completed_units += float(row["quantity"])
        elif row["status"] == "CANCELLED":
            cancelled_sales.add(sale_id)

    totals = list(completed_sales.values())
    total_sales_cents = sum(int(row["total_cents"]) for row in totals)
    total_profit_cents = sum(int(row["estimated_gross_profit_cents"]) for row in totals)

    summary.append([settings.PROJECT_NAME])
    summary.append(["Período exportado", period_label])
    export_time = generated_at or datetime.now(ZoneInfo(settings.BUSINESS_TIMEZONE))
    summary.append(["Generado el", export_time.replace(tzinfo=None)])
    summary["B3"].number_format = DATE_FORMAT
    summary.append(["Total vendido", round(total_sales_cents / 100, 2)])
    summary["B4"].number_format = CURRENCY_FORMAT
    summary.append(["Operaciones completadas", len(completed_sales)])
    summary.append(["Unidades vendidas", round(completed_units, 3)])
    summary.append(["Ganancia bruta estimada", round(total_profit_cents / 100, 2)])
    summary["B7"].number_format = CURRENCY_FORMAT
    summary.append(["Ventas anuladas", len(cancelled_sales)])
    summary["A1"].font = Font(size=16, bold=True, color="312E81")
    for row_index in range(2, summary.max_row + 1):
        summary.cell(row_index, 1).font = Font(bold=True)
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 28

    headers = [
        "Fecha y hora",
        "ID de venta",
        "N.º comprobante",
        "Código de producto",
        "Producto",
        "Cantidad",
        "Unidad",
        "Precio unitario",
        "Subtotal",
        "Costo unitario histórico",
        "Ganancia bruta estimada",
        "Estado",
        "Motivo de anulación",
    ]
    details.append(headers)
    for cell in details[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for row in rows:
        if row["product_id"] is None:
            continue

        detail_row = [
            _to_excel_datetime(row["sale_date"]),
            row["sale_id"],
            row["sale_number"],
            row["product_sku"],
            row["product_name"],
            float(row["quantity"]),
            row["unit_of_measure"],
            round(int(row["unit_price_cents"]) / 100, 2),
            round(int(row["subtotal_cents"]) / 100, 2),
            round(int(row["unit_cost_cents"]) / 100, 2),
            round(int(row["item_profit_cents"]) / 100, 2),
            "Completada" if row["status"] == "COMPLETED" else "Anulada",
            row["cancellation_reason"],
        ]
        details.append(detail_row)
        current_row = details.max_row
        details.cell(current_row, 1).number_format = DATE_FORMAT
        details.cell(current_row, 6).number_format = "0.000"
        for column in (8, 9, 10, 11):
            details.cell(current_row, column).number_format = CURRENCY_FORMAT
        for column in (2, 4, 5, 7, 12, 13):
            _write_text(details.cell(current_row, column), detail_row[column - 1])

    details.freeze_panes = "A2"
    details.auto_filter.ref = details.dimensions
    column_widths = [20, 38, 16, 20, 34, 14, 16, 19, 16, 24, 25, 16, 40]
    for index, width in enumerate(column_widths, start=1):
        details.column_dimensions[get_column_letter(index)].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
