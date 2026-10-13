from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

from app.services.sales_export_service import build_sales_workbook


def test_sales_workbook_summarizes_completed_sales_and_lists_cancellations():
    rows = [
        {
            "sale_id": "sale-completed",
            "sale_number": 1,
            "sale_date": "2026-10-13 10:00:00",
            "total_cents": 4500,
            "total_cost_cents": 3000,
            "estimated_gross_profit_cents": 1500,
            "status": "COMPLETED",
            "cancellation_reason": None,
            "product_id": "product-1",
            "product_sku": "RIBBON-1",
            "product_name": "Cinta",
            "quantity": 2.0,
            "unit_of_measure": "metro",
            "unit_price_cents": 2000,
            "subtotal_cents": 4000,
            "unit_cost_cents": 1500,
            "item_profit_cents": 1000,
        },
        {
            "sale_id": "sale-completed",
            "sale_number": 1,
            "sale_date": "2026-10-13 10:00:00",
            "total_cents": 4500,
            "total_cost_cents": 3000,
            "estimated_gross_profit_cents": 1500,
            "status": "COMPLETED",
            "cancellation_reason": None,
            "product_id": "product-2",
            "product_sku": "BUTTON-1",
            "product_name": "Botón",
            "quantity": 1.0,
            "unit_of_measure": "unidad",
            "unit_price_cents": 500,
            "subtotal_cents": 500,
            "unit_cost_cents": 0,
            "item_profit_cents": 500,
        },
        {
            "sale_id": "sale-cancelled",
            "sale_number": 2,
            "sale_date": "2026-10-14 11:00:00",
            "total_cents": 9900,
            "total_cost_cents": 6000,
            "estimated_gross_profit_cents": 3900,
            "status": "CANCELLED",
            "cancellation_reason": "=1+1",
            "product_id": "product-3",
            "product_sku": "+SKU",
            "product_name": "=HYPERLINK(\"https://example.com\")",
            "quantity": 3.0,
            "unit_of_measure": "unidad",
            "unit_price_cents": 3300,
            "subtotal_cents": 9900,
            "unit_cost_cents": 2000,
            "item_profit_cents": 3900,
        },
    ]

    content = build_sales_workbook(
        rows,
        year=2026,
        month=10,
        generated_at=datetime(2026, 10, 15, 12, 30),
    )
    workbook = load_workbook(BytesIO(content), data_only=False)
    summary = workbook["Resumen"]
    details = workbook["Detalle de ventas"]

    assert summary["B2"].value == "Octubre 2026"
    assert summary["B4"].value == 45
    assert summary["B5"].value == 1
    assert summary["B6"].value == 3
    assert summary["B7"].value == 15
    assert summary["B8"].value == 1
    assert details.max_row == 4
    assert details["L4"].value == "Anulada"
    assert details["E4"].value == '=HYPERLINK("https://example.com")'
    assert details["E4"].data_type == "s"
    assert details["D4"].value == "+SKU"
    assert details["D4"].data_type == "s"
    assert details["M4"].value == "=1+1"
    assert details["M4"].data_type == "s"


def test_empty_month_produces_valid_workbook_with_zero_totals():
    workbook = load_workbook(
        BytesIO(build_sales_workbook([], year=2026, month=2)),
        data_only=False,
    )

    assert workbook.sheetnames == ["Resumen", "Detalle de ventas"]
    assert workbook["Resumen"]["B4"].value == 0
    assert workbook["Resumen"]["B5"].value == 0
    assert workbook["Detalle de ventas"].max_row == 1
